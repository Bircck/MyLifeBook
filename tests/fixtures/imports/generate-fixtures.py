"""Rebuild the wholly fictional import fixture; requires python-docx.

The tiny PNG is original test geometry constructed from six RGB pixels.
No downloaded artwork, real biography, or private document is included.
"""
from pathlib import Path
import struct
import zlib
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent


def chunk(kind, data):
    return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)


png = (b'\x89PNG\r\n\x1a\n'
       + chunk(b'IHDR', struct.pack('!IIBBBBB', 3, 2, 8, 2, 0, 0, 0))
       + chunk(b'IDAT', zlib.compress(bytes([0, 94, 120, 112, 200, 176, 122, 146, 98, 71,
                                           0, 208, 210, 177, 75, 99, 107, 239, 227, 201])))
       + chunk(b'IEND', b''))
(HERE / 'detail.png').write_bytes(png)

doc = Document()
doc.core_properties.author = 'MyLifeBook fictional fixture'
doc.core_properties.title = 'The blue tin'
doc.core_properties.subject = 'Import fidelity test'
doc.sections[0].page_width = Inches(8.5)
doc.sections[0].page_height = Inches(11)
for style_name in ['Title', 'Heading 1', 'Heading 2', 'Normal']:
    style = doc.styles[style_name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.add_paragraph('The blue tin', 'Title')
doc.add_paragraph('This fictional source remembers an ordinary object and tests a portable import.')
doc.add_heading('A place for small things', 1)
paragraph = doc.add_paragraph('I kept ')
paragraph.add_run('buttons').bold = True
paragraph.add_run(' and ')
paragraph.add_run('bus tickets').italic = True
paragraph.add_run(' in a blue tin. This is invented sample writing.')
doc.add_paragraph('Keep the original wording', 'List Bullet')
doc.add_paragraph('Review every caption', 'List Bullet')
picture = doc.add_picture(str(HERE / 'detail.png'), width=Inches(2))
picture._inline.docPr.set('descr', 'Six colored squares used as an original test image')
doc.add_paragraph('Figure 1 A fictional keepsake illustration', 'Caption')


def link(paragraph, label, target):
    rel_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel_id)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


link(doc.add_paragraph(), 'Example source reference', 'https://example.org/source')
link(doc.add_paragraph(), 'Unsafe link label retained as text', 'javascript:alert(1)')
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
for row, values in zip(table.rows, [['Object', 'Meaning'], ['Tin', 'Small keepsakes']]):
    for cell, text in zip(row.cells, values):
        cell.text = text
doc.sections[0].footer.paragraphs[0].text = 'Fixture footer to compare with the preserved original'
doc.save(HERE / 'sample.docx')
print(HERE / 'sample.docx')
